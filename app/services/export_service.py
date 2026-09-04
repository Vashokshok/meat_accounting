import io
from datetime import date

from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.operation import Operation
from app.models.operation_change import OperationChange
from app.schemas.report import ReportOut


def _fill_sheet(ws, headers: list[str], rows: list[list]) -> None:
    """Шапка + строки в лист."""
    ws.append(headers)
    for row in rows:
        ws.append(row)


async def build_excel(
    db: AsyncSession, report: ReportOut, frm: date, to_excl: date
) -> io.BytesIO:
    """Excel из report_service (без своей математики)."""
    wb = Workbook()

    ws = wb.active
    ws.title = "Сводка"
    _fill_sheet(
        ws,
        [
            "Мясо",
            "На начало",
            "Приход",
            "Вертель",
            "Списание",
            "Конвектомат",
            "Франшиза",
            "На конец",
        ],
        [
            [
                i.meat_type,
                i.opening,
                i.incoming,
                i.spit,
                i.write_off,
                i.convection,
                i.franchise,
                i.closing,
            ]
            for i in report.items
        ],
    )

    ops = (
        await db.execute(
            select(Operation)
            .where(Operation.operation_date >= frm, Operation.operation_date < to_excl)
            .order_by(Operation.id)
        )
    ).scalars()
    ws2 = wb.create_sheet("Операции")
    _fill_sheet(
        ws2,
        ["ID", "Дата", "Тип", "Мясо", "Кг", "Франшиза", "Статус", "Автор", "Коммент"],
        [
            [
                o.id,
                o.operation_date.isoformat(),
                o.type,
                o.meat_type,
                float(o.quantity),
                o.franchise_id,
                o.status,
                o.created_by,
                o.comment,
            ]
            for o in ops
        ],
    )

    changes = (
        await db.execute(
            select(OperationChange).order_by(OperationChange.id).limit(5000)
        )
    ).scalars()
    ws3 = wb.create_sheet("История изменений")
    _fill_sheet(
        ws3,
        ["ID", "Операция", "Кто", "Когда", "Было", "Стало"],
        [
            [
                c.id,
                c.operation_id,
                c.changed_by,
                c.changed_at.isoformat(),
                str(c.old_data),
                str(c.new_data),
            ]
            for c in changes
        ],
    )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
